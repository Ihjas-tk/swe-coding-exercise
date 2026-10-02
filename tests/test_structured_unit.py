"""Structured loader on tiny generated XLSX/CSV files: header detection, typing, naming, guard, caps."""

import sqlite3

import openpyxl
import pytest

from ae.extract.structured import MAX_ROW_TEXT_COLS, SQLError, describe_schema, load_structured, query_sql


@pytest.fixture
def bom_xlsx(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Parts List"
    ws.append(["ACME Robot Arm — Bill of Materials"])  # title rows above the real header
    ws.append(["Revision B"])
    ws.append([])
    ws.append(["Part Number", "Description", "Qty", None])  # cost column has no header
    for i in range(1, 20):
        ws.append([f"P-{i:03d}", f"Widget {i}", i, 1.5 * i])
    ws.append(["P-020", "Widget 20", 20, "16H"])  # one text cell in a numeric column
    sup = wb.create_sheet("Suppliers")
    sup.append(["Supplier", "Country"])
    sup.append(["Acme", "US"])
    sup.append(["Bolt Co", "DE"])
    path = tmp_path / "bom.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def wide_csv(tmp_path):
    path = tmp_path / "wide.csv"
    header = ",".join(f"sensor_{i}" for i in range(60))
    rows = [",".join(str(i * k) for i in range(60)) for k in (1, 2)]
    path.write_text("\n".join([header, *rows]) + "\n")
    return path


@pytest.fixture
def db(tmp_path, bom_xlsx, wide_csv):
    path = tmp_path / "t.sqlite"
    load_structured([bom_xlsx, wide_csv], path)
    return path


def test_multi_sheet_workbook_gets_one_table_per_sheet(tmp_path, bom_xlsx):
    tables = load_structured([bom_xlsx], tmp_path / "m.sqlite")
    assert [(t.name, t.doc, t.n_rows) for t in tables] == [
        ("bom__parts_list", "bom.xlsx", 20),
        ("bom__suppliers", "bom.xlsx", 2),
    ]


def test_single_sheet_workbook_is_named_after_the_file(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["Name", "Value"])
    wb.active.append(["a", 1])
    path = tmp_path / "Single Sheet.xlsx"
    wb.save(path)
    (t,) = load_structured([path], tmp_path / "s.sqlite")
    assert t.name == "single_sheet"


def test_header_found_below_title_rows_and_unnamed_column_is_numbered(tmp_path, bom_xlsx):
    t = load_structured([bom_xlsx], tmp_path / "h.sqlite")[0]
    assert [(c.name, c.type) for c in t.columns] == [
        ("part_number", "text"),
        ("description", "text"),
        ("qty", "int"),
        ("col_4", "float"),
    ]


def test_provenance_row_numbers_start_at_one_after_the_header(db):
    assert query_sql("SELECT part_number FROM bom__parts_list WHERE _row = 1", db)["rows"] == [["P-001"]]


def test_numeric_majority_column_turns_text_cell_into_null(db):
    rows = query_sql("SELECT part_number, col_4 FROM bom__parts_list WHERE _row IN (19, 20) ORDER BY _row", db)["rows"]
    assert rows == [["P-019", 28.5], ["P-020", None]]
    assert query_sql("SELECT SUM(col_4), COUNT(col_4) FROM bom__parts_list", db)["rows"] == [[285.0, 19]]


def test_mostly_text_column_stays_text(tmp_path):
    path = tmp_path / "mixed.csv"
    path.write_text("id,code\n1,A1\n2,7\n3,B2\n")
    (t,) = load_structured([path], tmp_path / "x.sqlite")
    assert {c.name: c.type for c in t.columns} == {"id": "int", "code": "text"}


@pytest.mark.parametrize(
    "sql",
    [
        "DELETE FROM bom__suppliers",
        "SELECT 1; DROP TABLE bom__suppliers",
        "PRAGMA table_info(bom__suppliers)",
        "ATTACH DATABASE 'x.db' AS y",
        "UPDATE bom__suppliers SET country = 'FR'",
        "WITH x AS (SELECT 1) DELETE FROM bom__suppliers",
    ],
)
def test_read_only_guard_rejects_writes_and_multi_statements(db, sql):
    with pytest.raises(SQLError):
        query_sql(sql, db)
    assert query_sql("SELECT COUNT(*) FROM bom__suppliers", db)["rows"] == [[2]]


def test_read_only_guard_allows_select_cte_and_trailing_semicolon(db):
    assert query_sql("WITH a AS (SELECT 1 AS v) SELECT v FROM a", db)["rows"] == [[1]]
    assert query_sql("  select country from bom__suppliers order by country ;  ", db)["rows"] == [["DE"], ["US"]]


def test_read_only_guard_allows_keywords_inside_string_literals(db):
    assert query_sql("SELECT COUNT(*) FROM bom__suppliers WHERE country = 'firmware update'", db)["rows"] == [[0]]


def test_sql_errors_are_wrapped(db):
    with pytest.raises(SQLError, match="no such table"):
        query_sql("SELECT * FROM missing_table", db)


def test_query_results_are_truncated_at_limit(db):
    res = query_sql("SELECT part_number FROM bom__parts_list", db, limit=5)
    assert len(res["rows"]) == 5 and res["truncated"] is True and res["columns"] == ["part_number"]


def test_describe_schema_caps_wide_tables(db):
    schema = describe_schema(db)
    wide = schema.split('TABLE "wide"')[1]
    lines = wide.strip().splitlines()
    assert len(lines) == 1 + 20 + 1 + 8  # table line, first 20, summary, last 8
    assert "... 32 more columns named like the ones above (sensor_20 .. sensor_51), all INTEGER" in wide
    assert '"sensor_19" INTEGER' in wide and '"sensor_52" INTEGER' in wide and '"sensor_30"' not in wide


def test_describe_schema_lists_low_cardinality_text_values(db):
    assert '"country" TEXT  -- 2 distinct; values: DE, US' in describe_schema(db)


def test_row_text_is_capped_at_thirty_columns(db):
    conn = sqlite3.connect(db)
    (text,) = conn.execute("SELECT text FROM _structured_rows WHERE doc = 'wide.csv' AND row = 2").fetchone()
    conn.close()
    assert text.startswith("wide row 2: sensor_0=0; sensor_1=2;")
    assert text.count("=") == MAX_ROW_TEXT_COLS
    assert f"sensor_{MAX_ROW_TEXT_COLS - 1}=" in text and f"sensor_{MAX_ROW_TEXT_COLS}=" not in text
    assert text.endswith("… (30 more columns)")


def test_reloading_a_file_replaces_its_rows(tmp_path, bom_xlsx):
    path = tmp_path / "r.sqlite"
    load_structured([bom_xlsx], path)
    load_structured([bom_xlsx], path)
    assert query_sql("SELECT COUNT(*) FROM _structured_rows WHERE doc = 'bom.xlsx'", path)["rows"] == [[22]]
